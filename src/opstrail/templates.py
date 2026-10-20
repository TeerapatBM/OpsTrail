"""Document templates: templates/<name>/{template.toml, skeleton.md, guide.md, examples/, reference.docx}."""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

try:
    import tomllib  # type: ignore[import-not-found]
except ModuleNotFoundError:  # pragma: no cover
    import tomli as tomllib  # type: ignore[no-redef]

EXAMPLE_CHAR_LIMIT = 6000


@dataclass
class Template:
    name: str
    title: str
    description: str
    directory: Path
    skeleton: str
    guide: str
    title_format: str = "{title}"
    example_files: list[Path] = field(default_factory=list)

    @property
    def reference_docx(self) -> Path | None:
        path = self.directory / "reference.docx"
        return path if path.exists() else None

    def examples_text(self, max_examples: int = 2, limit: int = EXAMPLE_CHAR_LIMIT) -> list[tuple[str, str]]:
        """Return (filename, text) for up to `max_examples` readable examples."""
        out: list[tuple[str, str]] = []
        for path in self.example_files:
            if len(out) >= max_examples:
                break
            text = read_example(path)
            if text:
                if len(text) > limit:
                    text = text[:limit] + "\n…(truncated)"
                out.append((path.name, text))
        return out

    def summary(self) -> dict:
        return {
            "name": self.name,
            "title": self.title,
            "description": self.description,
            "examples": [p.name for p in self.example_files],
            "has_reference_docx": self.reference_docx is not None,
        }


def read_example(path: Path) -> str:
    suffix = path.suffix.lower()
    try:
        if suffix in (".md", ".txt", ".markdown"):
            return path.read_text(encoding="utf-8", errors="replace").strip()
        if suffix == ".docx":
            from docx import Document

            doc = Document(str(path))
            lines = []
            for para in doc.paragraphs:
                text = para.text.strip()
                if not text:
                    continue
                style = (para.style.name or "").lower() if para.style is not None else ""
                if style.startswith("heading"):
                    level = "".join(ch for ch in style if ch.isdigit()) or "1"
                    lines.append("#" * min(int(level), 6) + " " + text)
                else:
                    lines.append(text)
            for table in doc.tables:
                for row in table.rows:
                    lines.append("| " + " | ".join(c.text.strip() for c in row.cells) + " |")
            return "\n".join(lines).strip()
    except Exception:
        return ""
    return ""


def load_template(directory: Path) -> Template | None:
    if not directory.is_dir():
        return None
    meta: dict = {}
    meta_file = directory / "template.toml"
    if meta_file.exists():
        with meta_file.open("rb") as fh:
            meta = tomllib.load(fh)
    skeleton_file = directory / "skeleton.md"
    if not skeleton_file.exists():
        return None
    guide_file = directory / "guide.md"
    examples_dir = directory / "examples"
    examples = (
        sorted(p for p in examples_dir.iterdir() if p.is_file() and not p.name.startswith((".", "~$")))
        if examples_dir.is_dir()
        else []
    )
    return Template(
        name=directory.name,
        title=str(meta.get("title") or directory.name),
        description=str(meta.get("description") or ""),
        directory=directory,
        skeleton=skeleton_file.read_text(encoding="utf-8"),
        guide=guide_file.read_text(encoding="utf-8") if guide_file.exists() else "",
        title_format=str(meta.get("title_format") or "{title}"),
        example_files=examples,
    )


def load_templates(root: Path) -> dict[str, Template]:
    out: dict[str, Template] = {}
    if not root.is_dir():
        return out
    for child in sorted(root.iterdir()):
        if child.name.startswith((".", "_")):
            continue
        tpl = load_template(child)
        if tpl:
            out[tpl.name] = tpl
    return out

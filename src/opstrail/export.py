"""Export document markdown to .md (with assets) and .docx."""

from __future__ import annotations

import re
import shutil
from pathlib import Path
from typing import Callable

from opstrail.store import Store

SHOT_RE = re.compile(r"shot://(\d+)")
IMAGE_LINE_RE = re.compile(r"^\s*!\[([^\]]*)\]\(([^)\s]+)(?:\s+\"[^\"]*\")?\)\s*$")
INLINE_RE = re.compile(r"(\*\*[^*]+\*\*|__[^_]+__|`[^`]+`|\*[^*\s][^*]*\*|\[[^\]]+\]\([^)]+\))")
HR_RE = re.compile(r"^\s*(?:(?:-\s*){3,}|(?:\*\s*){3,}|(?:_\s*){3,})$")
INVALID_FS = re.compile(r'[<>:"/\\|?*\x00-\x1f]')


def slugify(text: str, limit: int = 60) -> str:
    text = INVALID_FS.sub("", text or "").strip()
    text = re.sub(r"\s+", "-", text)
    text = text.strip(".-") or "document"
    return text[:limit].rstrip(".-")


def shot_resolver(store: Store) -> Callable[[str], Path | None]:
    def resolve(ref: str) -> Path | None:
        m = SHOT_RE.fullmatch(ref.strip())
        if m:
            ev = store.get_event(int(m.group(1)))
            if ev and ev["type"] == "screenshot":
                path = Path(str(ev["data"].get("path", "")))
                return path if path.exists() else None
            return None
        path = Path(ref)
        return path if path.is_absolute() and path.exists() else None

    return resolve


# --------------------------------------------------------------------- markdown
def export_markdown(markdown: str, dest: Path, resolve: Callable[[str], Path | None]) -> Path:
    """Write markdown, copying referenced screenshots into <name>_assets/."""
    dest.parent.mkdir(parents=True, exist_ok=True)
    assets_dir = dest.with_name(dest.stem + "_assets")

    def replace(match: re.Match[str]) -> str:
        alt, ref = match.group(1), match.group(2)
        path = resolve(ref)
        if path is None:
            return f"![{alt}]({ref})" if not ref.startswith("shot://") else f"*[missing screenshot {ref}]*"
        assets_dir.mkdir(parents=True, exist_ok=True)
        target = assets_dir / path.name
        if not target.exists():
            shutil.copy2(path, target)
        return f"![{alt}]({assets_dir.name}/{target.name})"

    out = re.sub(r"!\[([^\]]*)\]\(([^)\s]+)\)", replace, markdown)
    dest.write_text(out, encoding="utf-8")
    return dest


# ------------------------------------------------------------------------- docx
def _set_fonts(style, font: str | None, size: float | None) -> None:
    from docx.oxml.ns import qn
    from docx.shared import Pt

    if font:
        style.font.name = font
        rpr = style.element.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = rpr.makeelement(qn("w:rFonts"), {})
            rpr.append(rfonts)
        for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            rfonts.set(qn(attr), font)
        for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            if rfonts.get(qn(attr)) is not None:
                del rfonts.attrib[qn(attr)]
    if size:
        style.font.size = Pt(size)
        rpr = style.element.get_or_add_rPr()
        szcs = rpr.find(qn("w:szCs"))
        if szcs is None:
            szcs = rpr.makeelement(qn("w:szCs"), {})
            rpr.append(szcs)
        szcs.set(qn("w:val"), str(int(round(size * 2))))


def _style(doc, name: str):
    try:
        return doc.styles[name]
    except KeyError:
        return None


def _add_inline(paragraph, text: str, mono_font: str = "Consolas") -> None:
    pos = 0
    for match in INLINE_RE.finditer(text):
        if match.start() > pos:
            paragraph.add_run(text[pos : match.start()])
        token = match.group(0)
        if token.startswith("**") or token.startswith("__"):
            paragraph.add_run(token[2:-2]).bold = True
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = mono_font
        elif token.startswith("["):
            m = re.match(r"\[([^\]]+)\]\(([^)]+)\)", token)
            if m:
                label, url = m.group(1), m.group(2)
                paragraph.add_run(label if label == url else f"{label} ({url})")
        else:
            paragraph.add_run(token[1:-1]).italic = True
        pos = match.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def _shade(paragraph, fill: str = "F2F2F2") -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    ppr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    ppr.append(shd)


def _clear_body(doc) -> None:
    body = doc.element.body
    for child in list(body):
        if not child.tag.endswith("}sectPr"):
            body.remove(child)


def _split_row(line: str) -> list[str]:
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    return [cell.strip() for cell in re.split(r"(?<!\\)\|", line)]


def markdown_to_docx(
    markdown: str,
    dest: Path,
    resolve: Callable[[str], Path | None],
    *,
    reference: Path | None = None,
    font: str | None = "Leelawadee UI",
    font_size: float | None = 10.5,
    image_width_in: float = 6.0,
) -> Path:
    from docx import Document
    from docx.shared import Inches, Pt

    doc = Document(str(reference)) if reference else Document()
    if reference:
        _clear_body(doc)
    else:
        for name in ("Normal", "List Bullet", "List Bullet 2", "Quote", "Table Grid"):
            st = _style(doc, name)
            if st is not None:
                _set_fonts(st, font, font_size if name == "Normal" else None)
        for name in ("Title", "Heading 1", "Heading 2", "Heading 3", "Heading 4"):
            st = _style(doc, name)
            if st is not None:
                _set_fonts(st, font, None)

    lines = markdown.replace("\r\n", "\n").split("\n")
    h1_count = sum(1 for ln in lines if re.match(r"^#\s+", ln))
    shift = 1 if h1_count == 1 else 0
    seen_title = False
    para_buffer: list[str] = []

    def flush_para() -> None:
        if para_buffer:
            p = doc.add_paragraph()
            _add_inline(p, " ".join(s.strip() for s in para_buffer))
            para_buffer.clear()

    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # fenced code
        if stripped.startswith("```"):
            flush_para()
            code: list[str] = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code.append(lines[i])
                i += 1
            i += 1  # closing fence
            p = doc.add_paragraph()
            _shade(p)
            for n, code_line in enumerate(code):
                run = p.add_run(code_line)
                run.font.name = "Consolas"
                run.font.size = Pt(9)
                if n < len(code) - 1:
                    run.add_break()
            continue

        # heading
        m = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if m:
            flush_para()
            level = len(m.group(1))
            text = m.group(2).strip().rstrip("#").strip()
            if level == 1 and shift == 1 and not seen_title:
                doc.add_heading(text, level=0)
                seen_title = True
            else:
                doc.add_heading(text, level=max(1, min(level - shift, 9)))
            i += 1
            continue

        # table
        if stripped.startswith("|") and i + 1 < len(lines) and re.match(r"^\s*\|?\s*:?-{2,}", lines[i + 1]):
            flush_para()
            header = _split_row(stripped)
            rows: list[list[str]] = []
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(_split_row(lines[i]))
                i += 1
            cols = max([len(header)] + [len(r) for r in rows])
            table = doc.add_table(rows=1 + len(rows), cols=cols)
            if _style(doc, "Table Grid") is not None:
                table.style = doc.styles["Table Grid"]
            for c in range(cols):
                cell_p = table.rows[0].cells[c].paragraphs[0]
                _add_inline(cell_p, header[c] if c < len(header) else "")
                for run in cell_p.runs:
                    run.bold = True
            for r, row in enumerate(rows, start=1):
                for c in range(cols):
                    _add_inline(table.rows[r].cells[c].paragraphs[0], row[c] if c < len(row) else "")
            doc.add_paragraph()
            continue

        # image on its own line
        m = IMAGE_LINE_RE.match(line)
        if m:
            flush_para()
            alt, ref = m.group(1), m.group(2)
            path = resolve(ref)
            if path is not None:
                try:
                    doc.add_picture(str(path), width=Inches(image_width_in))
                except Exception:
                    doc.add_paragraph(f"[image could not be embedded: {ref}]")
            else:
                doc.add_paragraph(f"[missing image: {ref}]")
            if alt:
                cap = doc.add_paragraph()
                cap.add_run(alt).italic = True
            i += 1
            continue

        # bullets / checkboxes
        m = re.match(r"^(\s*)[-*+]\s+(.*)$", line)
        if m and not HR_RE.match(line):
            flush_para()
            indent = len(m.group(1).replace("\t", "  "))
            text = m.group(2)
            text = re.sub(r"^\[ \]\s*", "☐ ", text)
            text = re.sub(r"^\[[xX]\]\s*", "☑ ", text)
            style_name = "List Bullet 2" if indent >= 2 and _style(doc, "List Bullet 2") is not None else "List Bullet"
            p = doc.add_paragraph(style=style_name if _style(doc, style_name) is not None else None)
            _add_inline(p, text)
            i += 1
            continue

        # numbered list (literal numbers – avoids Word's shared numbering bug)
        m = re.match(r"^(\s*)(\d+)[.)]\s+(.*)$", line)
        if m:
            flush_para()
            indent = len(m.group(1).replace("\t", "  "))
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3 + 0.25 * (indent // 2))
            p.paragraph_format.first_line_indent = Inches(-0.25)
            _add_inline(p, f"{m.group(2)}. {m.group(3)}")
            i += 1
            continue

        # blockquote
        m = re.match(r"^\s*>\s?(.*)$", line)
        if m:
            flush_para()
            p = doc.add_paragraph(style="Quote" if _style(doc, "Quote") is not None else None)
            _add_inline(p, m.group(1))
            i += 1
            continue

        # horizontal rule
        if HR_RE.match(line):
            flush_para()
            i += 1
            continue

        if not stripped:
            flush_para()
            i += 1
            continue

        para_buffer.append(line)
        i += 1

    flush_para()
    dest.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(dest))
    return dest

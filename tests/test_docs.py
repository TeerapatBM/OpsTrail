import pytest
from conftest import DAY
from docx import Document

from opstrail.service import OpsTrailError


def test_templates_load(svc):
    names = {t["name"] for t in svc.list_templates()}
    assert {"sop", "delivery", "guide", "daily-worklog"} <= names


def test_doc_session_flow_and_export(svc, sample_day, cfg):
    svc.refresh()
    ep = svc.list_episodes(DAY)[0]
    res = svc.start_doc_session("sop", "ติดตั้ง Splunk Universal Forwarder", [ep["id"]])
    sid = res["session_id"]
    pack = res["source_pack"]
    assert "Writing rules" in pack and "Skeleton" in pack
    assert "shot://" in pack and "[REDACTED]" in pack and "Sup3rS3cret" not in pack
    assert "ใช้ port 9997" in pack  # user note is in the pack
    with pytest.raises(OpsTrailError):
        svc.export_doc(sid)  # nothing saved yet

    shot_id = next(e["id"] for e in svc.store.events_for_date(DAY) if e["type"] == "screenshot")
    md = f"""# SOP: ติดตั้ง Splunk UF

| รายการ | รายละเอียด |
|---|---|
| Version | 1.0 |

## 1. วัตถุประสงค์

ติดตั้ง **Universal Forwarder** และส่ง log ไป `10.0.0.5:9997`

## 5. ขั้นตอน

1. ติดตั้ง package
2. Start service

- [ ] ตรวจสอบ
- bullet with [link](https://docs.splunk.com)

```bash
sudo dpkg -i splunkforwarder.deb
/opt/splunkforwarder/bin/splunk start --password [REDACTED]
```

![UF connected](shot://{shot_id})

> หมายเหตุ: ใช้ port ตาม design

---
"""
    saved = svc.save_draft(sid, md, "first draft")
    assert saved["version"] == 1
    assert svc.save_draft(sid, md + "\nmore", "v2")["version"] == 2

    out = svc.export_doc(sid, "both")
    md_path = next(p for p in out["files"] if p.endswith(".md"))
    docx_path = next(p for p in out["files"] if p.endswith(".docx"))
    text = open(md_path, encoding="utf-8").read()
    assert "_assets/094000.png" in text
    doc = Document(docx_path)
    body = "\n".join(p.text for p in doc.paragraphs)
    assert "ติดตั้ง" in body and "1. ติดตั้ง package" in body and "☐ ตรวจสอบ" in body
    assert doc.paragraphs[0].style.name == "Title"
    assert len(doc.inline_shapes) == 1
    assert len(doc.tables) == 1

    sess = svc.get_doc_session(sid)
    assert sess["latest_markdown"].endswith("more") and len(sess["versions"]) == 2

    assert svc.search_docs("Universal Forwarder")[0]["session_id"] == sid
    assert svc.search_docs("ติดตั้ง")  # Thai substring search
    assert svc.set_doc_status(sid, "review")["status"] == "review"
    with pytest.raises(OpsTrailError):
        svc.set_doc_status(sid, "bogus")


def test_continue_session_with_new_sources(svc, sample_day):
    svc.refresh()
    a, b = svc.list_episodes(DAY)
    sid = svc.start_doc_session("delivery", "Customer X delivery", [a["id"]])["session_id"]
    res = svc.add_doc_sources(sid, [b["id"]])
    assert res["sources_added"] == [f"episode #{b['id']}"]
    assert "Sentinel" in res["new_material"] and "Splunk Docs" not in res["new_material"]
    again = svc.add_doc_sources(sid, [b["id"]])
    assert again["sources_added"] == [] and again["new_material"] == "_nothing new_"


def test_date_source_and_validation(svc, sample_day):
    svc.refresh()
    res = svc.start_doc_session("guide", "KQL basics", date=DAY)
    assert "Episode #" in res["source_pack"]
    with pytest.raises(OpsTrailError):
        svc.start_doc_session("nope", "x", date=DAY)
    with pytest.raises(OpsTrailError):
        svc.start_doc_session("sop", "x")
    with pytest.raises(OpsTrailError):
        svc.list_episodes("2026-13-01")


def test_daily_recap(svc, sample_day):
    res = svc.daily_recap(DAY)
    md = res["markdown"]
    assert md.startswith(f"# Daily Worklog — {DAY}")
    assert "Doc candidates" in md and "ใช้ port 9997" in md
    again = svc.daily_recap(DAY)
    assert again["session_id"] == res["session_id"] and again["version"] == 2
    assert svc.export_doc(res["session_id"], "docx")["files"][0].endswith(".docx")


def test_reference_docx_styles(svc, sample_day, cfg):
    tpl_dir = cfg.templates_dir / "sop"
    ref = tpl_dir / "reference.docx"
    created = not ref.exists()
    if created:
        d = Document()
        d.add_paragraph("company cover text that must not appear")
        d.save(str(ref))
    try:
        svc.refresh()
        ep = svc.list_episodes(DAY)[0]
        sid = svc.start_doc_session("sop", "With reference", [ep["id"]])["session_id"]
        svc.save_draft(sid, "# Title\n\nHello")
        path = svc.export_doc(sid, "docx")["files"][0]
        body = [p.text for p in Document(path).paragraphs]
        assert "company cover text that must not appear" not in body and "Hello" in body
    finally:
        if created:
            ref.unlink()


def test_purge_keeps_documents(svc, sample_day):
    svc.refresh()
    ep = svc.list_episodes(DAY)[0]
    sid = svc.start_doc_session("sop", "keep me", [ep["id"]])["session_id"]
    svc.save_draft(sid, "# keep")
    assert svc.purge(date=DAY) > 0
    assert svc.store.events_for_date(DAY) == []
    assert svc.get_doc_session(sid)["latest_markdown"] == "# keep"
